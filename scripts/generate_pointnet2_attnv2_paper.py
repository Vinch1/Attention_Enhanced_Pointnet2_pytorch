#!/usr/bin/env python3
from __future__ import annotations

import csv
import importlib
import json
import re
import subprocess
import sys
import textwrap
from collections import Counter
from pathlib import Path

import h5py
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch

try:
    from docx import Document
    from docx.enum.section import WD_SECTION_START
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
except ModuleNotFoundError as exc:
    raise SystemExit(
        "python-docx is required. Run with: uv run --with python-docx scripts/generate_pointnet2_attnv2_paper.py"
    ) from exc


ROOT = Path(__file__).resolve().parents[1]
PAPER_DIR = ROOT / "docs" / "papers"
ASSET_DIR = PAPER_DIR / "assets" / "pointnet2_attnv2_single_column"
DOCX_PATH = PAPER_DIR / "2026-04-14-pointnet2-attnv2-single-column-paper.docx"

MODEL_PATHS = {
    "baseline": "pointnet2_cls_ssg",
    "se": "pointnet2_cls_ssg_se",
    "attn_v1": "pointnet2_cls_ssg_attn",
    "attn_v2": "pointnet2_cls_ssg_attn_v2",
}

MODEL_LABELS = {
    "baseline": "SSG Baseline",
    "se": "SSG + SE",
    "attn_v1": "Attention v1",
    "attn_v2": "Attention v2",
}

MODEL_COLORS = {
    "baseline": "#12304B",
    "se": "#1D6FD6",
    "attn_v1": "#D66C3B",
    "attn_v2": "#0F9960",
}

REFERENCE_TEXT = [
    "[1] C. R. Qi, H. Su, K. Mo, and L. J. Guibas, 'PointNet: Deep Learning on Point Sets for 3D Classification and Segmentation,' in Proceedings of the IEEE Conference on Computer Vision and Pattern Recognition, 2017.",
    "[2] C. R. Qi, L. Yi, H. Su, and L. J. Guibas, 'PointNet++: Deep Hierarchical Feature Learning on Point Sets in a Metric Space,' in Advances in Neural Information Processing Systems, 2017.",
    "[3] J. Hu, L. Shen, and G. Sun, 'Squeeze-and-Excitation Networks,' in Proceedings of the IEEE Conference on Computer Vision and Pattern Recognition, 2018.",
    "[4] Y. Wang, Y. Sun, Z. Liu, S. E. Sarma, M. M. Bronstein, and J. M. Solomon, 'Dynamic Graph CNN for Learning on Point Clouds,' ACM Transactions on Graphics, vol. 38, no. 5, 2019.",
    "[5] J. Lee, Y. Lee, J. Kim, A. Kosiorek, S. Choi, and Y. W. Teh, 'Set Transformer: A Framework for Attention-based Permutation-Invariant Neural Networks,' in Proceedings of the 36th International Conference on Machine Learning, 2019.",
    "[6] Y. Li, R. Bu, M. Sun, and B. Chen, 'PointCNN,' in Advances in Neural Information Processing Systems, 2018.",
    "[7] W. Wu, Z. Qi, and L. Fuxin, 'PointConv: Deep Convolutional Networks on 3D Point Clouds,' in Proceedings of the IEEE Conference on Computer Vision and Pattern Recognition, 2019.",
    "[8] H. Zhao, L. Jiang, J. Jia, P. Torr, and V. Koltun, 'Point Transformer,' in Proceedings of the IEEE/CVF International Conference on Computer Vision, 2021.",
    "[9] M. A. Uy, Q. T. Pham, B.-S. Hua, T. Nguyen, and S.-K. Yeung, 'Revisiting Point Cloud Classification: A New Benchmark Dataset and Classification Model on Real-World Data,' in Proceedings of the IEEE/CVF International Conference on Computer Vision, 2019.",
    "[10] A. Vaswani et al., 'Attention Is All You Need,' in Advances in Neural Information Processing Systems, 2017.",
]


def pct(value: float, digits: int = 3) -> str:
    return f"{value * 100:.{digits}f}%"


def million(value: int) -> str:
    return f"{value / 1_000_000:.2f}M"


def ensure_dirs() -> None:
    PAPER_DIR.mkdir(parents=True, exist_ok=True)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)


def parse_modelnet_log(path: Path) -> dict[str, float | int]:
    epoch_re = re.compile(r"Epoch \d+ \((\d+)/(\d+)\):")
    train_re = re.compile(r"Train Instance Accuracy: ([0-9.]+)")
    test_re = re.compile(r"Test Instance Accuracy: ([0-9.]+), Class Accuracy: ([0-9.]+)")
    batch_re = re.compile(r"batch_size=(\d+)")
    current_epoch = None
    batch_size = None
    best_instance = -1.0
    best_class = -1.0
    best_save_epoch = None
    final_train = None
    final_test = None
    for line in path.read_text().splitlines():
        batch_match = batch_re.search(line)
        if batch_match:
            batch_size = int(batch_match.group(1))
        epoch_match = epoch_re.search(line)
        if epoch_match:
            current_epoch = int(epoch_match.group(1))
        train_match = train_re.search(line)
        if train_match:
            final_train = float(train_match.group(1))
        test_match = test_re.search(line)
        if test_match:
            instance = float(test_match.group(1))
            class_acc = float(test_match.group(2))
            final_test = (instance, class_acc)
            best_instance = max(best_instance, instance)
            best_class = max(best_class, class_acc)
        if "Saving at log/classification" in line:
            best_save_epoch = current_epoch
    return {
        "batch_size": batch_size or 0,
        "best_instance": best_instance,
        "best_class": best_class,
        "best_epoch": best_save_epoch or 0,
        "final_train": final_train or 0.0,
        "final_test_instance": final_test[0] if final_test else 0.0,
        "final_test_class": final_test[1] if final_test else 0.0,
    }


def load_metrics(path: Path) -> list[dict[str, float]]:
    rows: list[dict[str, float]] = []
    with path.open() as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            rows.append({key: float(value) for key, value in row.items()})
    return rows


def metrics_snapshot(metrics: list[dict[str, float]], epochs: list[int]) -> list[list[str]]:
    lookup = {int(row["epoch"]): row for row in metrics}
    rows: list[list[str]] = []
    for epoch in epochs:
        row = lookup[epoch]
        rows.append(
            [
                str(epoch),
                pct(row["train_instance_acc"]),
                pct(row["test_instance_acc"]),
                pct(row["test_class_acc"]),
            ]
        )
    return rows


def load_run_summary(path: Path) -> dict[str, object]:
    return json.loads(path.read_text())


def dataset_stats() -> dict[str, dict[str, object]]:
    stats: dict[str, dict[str, object]] = {}
    for split, file_name in [("train", "train_point_clouds.h5"), ("test", "test_point_clouds.h5")]:
        counts: list[int] = []
        labels: list[int] = []
        with h5py.File(ROOT / "3D MNIST" / file_name, "r") as handle:
            for key in handle.keys():
                group = handle[key]
                counts.append(int(group["points"].shape[0]))
                labels.append(int(group.attrs["label"]))
        stats[split] = {
            "n": len(counts),
            "min": min(counts),
            "max": max(counts),
            "avg": round(sum(counts) / len(counts), 2),
            "counts": counts,
            "dist": dict(sorted(Counter(labels).items())),
        }
    return stats


def parameter_counts() -> tuple[dict[str, int], list[tuple[str, int]]]:
    sys.path.insert(0, str(ROOT / "models"))
    counts: dict[str, int] = {}
    top_params: list[tuple[str, int]] = []
    for key, module_name in MODEL_PATHS.items():
        module = importlib.import_module(module_name)
        model = module.get_model(40, normal_channel=False)
        counts[key] = sum(parameter.numel() for parameter in model.parameters())
        if key == "attn_v2":
            top_params = sorted(
                ((name, parameter.numel()) for name, parameter in model.named_parameters()),
                key=lambda item: item[1],
                reverse=True,
            )[:10]
    return counts, top_params


def git_overview() -> dict[str, list[str] | str]:
    branches = subprocess.check_output(["git", "branch", "--list"], cwd=ROOT, text=True).splitlines()
    log_lines = subprocess.check_output(
        ["git", "log", "--oneline", "--decorate", "--all", "--max-count=12"],
        cwd=ROOT,
        text=True,
    ).splitlines()
    return {
        "branches": [line.strip().lstrip("* ").strip() for line in branches],
        "log": log_lines,
    }


def set_matplotlib_defaults() -> None:
    plt.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "figure.facecolor": "white",
            "axes.facecolor": "white",
            "axes.spines.top": False,
            "axes.spines.right": False,
            "axes.edgecolor": "#7D8A98",
            "axes.labelcolor": "#16324F",
            "xtick.color": "#16324F",
            "ytick.color": "#16324F",
            "grid.color": "#E5EBF1",
            "grid.linewidth": 0.8,
        }
    )


def save_timeline_figure(path: Path) -> None:
    set_matplotlib_defaults()
    fig, ax = plt.subplots(figsize=(13, 3.4), dpi=220)
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 3)
    ax.axis("off")
    ax.plot([0.8, 9.2], [1.5, 1.5], color="#D7E1EA", linewidth=4, zorder=0)
    stages = [
        (1.2, "master", "Upstream\nbaseline", "#12304B", "eb64fe0"),
        (3.4, "feat/apple_silicon", "MPS safety,\nSE, early attn", "#1D6FD6", "2284629 -> 3a536bf"),
        (5.9, "pointnet2-attn-v2-mac-safe", "Masked hybrid\nattention repair", "#0F9960", "170cabf"),
        (8.4, "3dmnist-attnv2-training", "3D MNIST +\nstructured logging", "#7A4CC2", "c51fe93 -> HEAD"),
    ]
    for x, label, subtitle, color, commit in stages:
        ax.scatter([x], [1.5], s=620, color=color, edgecolor="white", linewidth=2.5, zorder=2)
        card = FancyBboxPatch(
            (x - 0.95, 1.92),
            1.9,
            0.78,
            boxstyle="round,pad=0.03,rounding_size=0.09",
            linewidth=1.1,
            edgecolor="#D6E0EA",
            facecolor="#F7FAFC",
        )
        ax.add_patch(card)
        ax.text(x, 2.48, label, ha="center", va="center", fontsize=10.5, fontweight="bold", color="#12304B")
        ax.text(x, 2.14, subtitle, ha="center", va="center", fontsize=9, color="#425466")
        ax.text(x, 1.03, commit, ha="center", va="center", fontsize=9, color=color, fontweight="bold")
    ax.text(0.82, 2.92, "Repository and experiment evolution", fontsize=14, fontweight="bold", color="#12304B")
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)


def save_repair_figure(path: Path) -> None:
    set_matplotlib_defaults()
    fig, ax = plt.subplots(figsize=(11.6, 7.2), dpi=220)
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")
    ax.text(0.22, 0.94, "Attention v1 failure modes", fontsize=15, fontweight="bold", color="#A8442B", ha="center", transform=ax.transAxes)
    ax.text(0.78, 0.94, "Attention v2 repair actions", fontsize=15, fontweight="bold", color="#0F6D4B", ha="center", transform=ax.transAxes)
    failures = [
        ("Padding bias", "Duplicated padded neighbors are harmless for max pooling but distort softmax attention by absorbing real probability mass."),
        ("Over-compressed tokens", "Reducing each hierarchy to one mean token discards too much multi-scale structure before fusion."),
        ("Weak classifier", "The fusion output is sent to an almost linear head, which limits nonlinear class separation."),
        ("MPS fragility", "Index writes, infinite masking values, and layout assumptions are more brittle on Apple Silicon."),
    ]
    repairs = [
        ("Explicit masks", "Attention is normalized only over real neighbors, while padded slots are kept only for safe indexing."),
        ("Hybrid pooling", "Attention pooling is fused with a max-pooling residual path to preserve PointNet++ robustness."),
        ("Richer fusion", "Mean and max summaries, a CLS token, and a global residual preserve more cross-level signal."),
        ("Numerical safeguards", "Finite masking values, nan cleanup, clamped indices, and contiguous tensors stabilize execution."),
    ]
    y_positions = [0.70, 0.49, 0.28, 0.07]
    for (left_title, left_body), (right_title, right_body), y in zip(failures, repairs, y_positions):
        left_box = FancyBboxPatch(
            (0.08, y),
            0.23,
            0.15,
            boxstyle="round,pad=0.04,rounding_size=0.08",
            linewidth=1.1,
            edgecolor="#E6C3B8",
            facecolor="#FFF6F2",
            transform=ax.transAxes,
        )
        right_box = FancyBboxPatch(
            (0.69, y),
            0.23,
            0.15,
            boxstyle="round,pad=0.04,rounding_size=0.08",
            linewidth=1.1,
            edgecolor="#B7D9C9",
            facecolor="#F3FBF7",
            transform=ax.transAxes,
        )
        ax.add_patch(left_box)
        ax.add_patch(right_box)
        ax.text(0.105, y + 0.112, textwrap.fill(left_title, width=18), fontsize=11.2, fontweight="bold", color="#A8442B", transform=ax.transAxes)
        ax.text(0.105, y + 0.035, textwrap.fill(left_body, width=24), fontsize=8.8, color="#4C5661", va="center", transform=ax.transAxes)
        ax.text(0.715, y + 0.112, textwrap.fill(right_title, width=18), fontsize=11.2, fontweight="bold", color="#0F6D4B", transform=ax.transAxes)
        ax.text(0.715, y + 0.035, textwrap.fill(right_body, width=24), fontsize=8.8, color="#4C5661", va="center", transform=ax.transAxes)
        arrow = FancyArrowPatch((0.37, y + 0.075), (0.63, y + 0.075), arrowstyle="simple", mutation_scale=18, color="#9BB7D0", transform=ax.transAxes)
        ax.add_patch(arrow)
    ax.text(
        0.5,
        0.03,
        "The key methodological contribution is not attention alone, but the diagnosis-and-repair loop that made the attention stack compatible with PointNet++ grouping semantics and MPS execution.",
        fontsize=10.3,
        color="#425466",
        ha="center",
        transform=ax.transAxes,
    )
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)


def save_modelnet_figure(path: Path, results: dict[str, dict[str, float | int]]) -> None:
    set_matplotlib_defaults()
    fig, ax = plt.subplots(figsize=(11.8, 5.6), dpi=220)
    labels = list(MODEL_PATHS.keys())
    x = np.arange(len(labels))
    width = 0.34
    best_instance = [results[key]["best_instance"] * 100 for key in labels]
    best_class = [results[key]["best_class"] * 100 for key in labels]
    bars1 = ax.bar(x - width / 2, best_instance, width, color=[MODEL_COLORS[key] for key in labels], label="Best instance accuracy")
    bars2 = ax.bar(x + width / 2, best_class, width, color="#8FB8DE", label="Best class accuracy")
    ax.set_xticks(x, [MODEL_LABELS[key] for key in labels])
    ax.set_ylim(84.0, 93.2)
    ax.set_ylabel("Accuracy (%)")
    ax.set_title("ModelNet40 comparison from raw text logs", fontsize=15, fontweight="bold", color="#12304B")
    ax.grid(axis="y", linestyle="--", alpha=0.7)
    for bar in list(bars1) + list(bars2):
        ax.text(bar.get_x() + bar.get_width() / 2, bar.get_height() + 0.08, f"{bar.get_height():.2f}", ha="center", va="bottom", fontsize=9)
    ax.legend(frameon=False, loc="upper left")
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)


def save_3dmnist_curves(path: Path, baseline_metrics: list[dict[str, float]], attn_metrics: list[dict[str, float]]) -> None:
    set_matplotlib_defaults()
    fig, axes = plt.subplots(1, 2, figsize=(12.6, 5.3), dpi=220, sharex=True)
    epochs = [row["epoch"] for row in baseline_metrics]
    axes[0].plot(epochs, [row["test_instance_acc"] * 100 for row in baseline_metrics], color=MODEL_COLORS["baseline"], linewidth=2.5, label="SSG baseline")
    axes[0].plot(epochs, [row["test_instance_acc"] * 100 for row in attn_metrics], color=MODEL_COLORS["attn_v2"], linewidth=2.5, label="Attention v2")
    axes[0].set_title("3D MNIST test instance accuracy", fontsize=14, fontweight="bold", color="#12304B")
    axes[0].set_ylabel("Accuracy (%)")
    axes[0].set_xlabel("Epoch")
    axes[0].set_ylim(30, 100)
    axes[0].grid(True, linestyle="--", alpha=0.7)
    axes[0].legend(frameon=False, loc="lower right")

    axes[1].plot(epochs, [row["test_class_acc"] * 100 for row in baseline_metrics], color=MODEL_COLORS["baseline"], linewidth=2.5, label="SSG baseline")
    axes[1].plot(epochs, [row["test_class_acc"] * 100 for row in attn_metrics], color=MODEL_COLORS["attn_v2"], linewidth=2.5, label="Attention v2")
    axes[1].set_title("3D MNIST test class accuracy", fontsize=14, fontweight="bold", color="#12304B")
    axes[1].set_xlabel("Epoch")
    axes[1].set_ylim(30, 100)
    axes[1].grid(True, linestyle="--", alpha=0.7)
    for axis in axes:
        axis.spines["left"].set_color("#7D8A98")
        axis.spines["bottom"].set_color("#7D8A98")
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)


def save_efficiency_figure(path: Path, params: dict[str, int], modelnet: dict[str, dict[str, float | int]]) -> None:
    set_matplotlib_defaults()
    fig, axes = plt.subplots(1, 2, figsize=(12.2, 5.3), dpi=220)
    labels = list(MODEL_PATHS.keys())
    param_values = [params[key] / 1_000_000 for key in labels]
    axes[0].bar(
        np.arange(len(labels)),
        param_values,
        color=[MODEL_COLORS[key] for key in labels],
        width=0.6,
    )
    axes[0].set_xticks(np.arange(len(labels)), [MODEL_LABELS[key] for key in labels], rotation=10)
    axes[0].set_ylabel("Parameters (millions)")
    axes[0].set_title("Model size expansion", fontsize=14, fontweight="bold", color="#12304B")
    axes[0].grid(axis="y", linestyle="--", alpha=0.7)
    for index, value in enumerate(param_values):
        axes[0].text(index, value + 0.05, f"{value:.2f}", ha="center", va="bottom", fontsize=9)

    for key in labels:
        axes[1].scatter(
            params[key] / 1_000_000,
            modelnet[key]["best_instance"] * 100,
            s=220,
            color=MODEL_COLORS[key],
            edgecolor="white",
            linewidth=1.4,
        )
        axes[1].annotate(
            MODEL_LABELS[key],
            (params[key] / 1_000_000, modelnet[key]["best_instance"] * 100),
            textcoords="offset points",
            xytext=(8, 6),
            fontsize=9.2,
        )
    axes[1].set_xlabel("Parameters (millions)")
    axes[1].set_ylabel("Best ModelNet40 instance accuracy (%)")
    axes[1].set_title("Accuracy versus parameter cost", fontsize=14, fontweight="bold", color="#12304B")
    axes[1].set_xlim(1.1, 6.3)
    axes[1].set_ylim(91.0, 92.7)
    axes[1].grid(True, linestyle="--", alpha=0.7)
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)


def save_dataset_figure(path: Path, stats: dict[str, dict[str, object]]) -> None:
    set_matplotlib_defaults()
    fig, axes = plt.subplots(1, 2, figsize=(12.6, 5.3), dpi=220)
    train_counts = np.array(stats["train"]["counts"], dtype=float)
    test_counts = np.array(stats["test"]["counts"], dtype=float)
    axes[0].hist(train_counts, bins=20, color="#1D6FD6", alpha=0.72, label="Train")
    axes[0].hist(test_counts, bins=20, color="#0F9960", alpha=0.6, label="Test")
    axes[0].set_title("3D MNIST point-count distribution", fontsize=14, fontweight="bold", color="#12304B")
    axes[0].set_xlabel("Raw points per sample")
    axes[0].set_ylabel("Frequency")
    axes[0].legend(frameon=False)
    axes[0].grid(axis="y", linestyle="--", alpha=0.7)

    labels = list(stats["train"]["dist"].keys())
    train_dist = [stats["train"]["dist"][label] for label in labels]
    test_dist = [stats["test"]["dist"][label] for label in labels]
    x = np.arange(len(labels))
    width = 0.38
    axes[1].bar(x - width / 2, train_dist, width, color="#1D6FD6", label="Train")
    axes[1].bar(x + width / 2, test_dist, width, color="#0F9960", label="Test")
    axes[1].set_xticks(x, [str(label) for label in labels])
    axes[1].set_title("3D MNIST class distribution", fontsize=14, fontweight="bold", color="#12304B")
    axes[1].set_xlabel("Digit label")
    axes[1].set_ylabel("Samples")
    axes[1].grid(axis="y", linestyle="--", alpha=0.7)
    axes[1].legend(frameon=False)
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)


def shade_cell(cell, fill: str) -> None:
    cell_properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    cell_properties.append(shading)


def set_cell_text(cell, text: str, *, bold: bool = False, color: str = "1F2933", align: int = WD_ALIGN_PARAGRAPH.LEFT) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_end])


def style_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    heading1 = document.styles["Heading 1"]
    heading1.font.name = "Cambria"
    heading1.font.size = Pt(16)
    heading1.font.bold = True
    heading1.font.color.rgb = RGBColor.from_string("12304B")

    heading2 = document.styles["Heading 2"]
    heading2.font.name = "Cambria"
    heading2.font.size = Pt(13)
    heading2.font.bold = True
    heading2.font.color.rgb = RGBColor.from_string("12304B")

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header.add_run("Repairing Attention for PointNet++ Classification")
    header_run.font.name = "Cambria"
    header_run.font.size = Pt(9)
    header_run.font.color.rgb = RGBColor.from_string("5B6B7A")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(footer)


def add_title_block(document: Document) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Repairing Attention for PointNet++ Classification")
    run.font.name = "Cambria"
    run.font.size = Pt(22)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string("12304B")

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("A Single-Column Technical Study Across ModelNet40 and 3D MNIST")
    run.font.name = "Cambria"
    run.font.size = Pt(13)
    run.italic = True
    run.font.color.rgb = RGBColor.from_string("2E6171")

    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run("Prepared on April 14, 2026")
    run.font.name = "Times New Roman"
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor.from_string("5B6B7A")

    rule = document.add_paragraph()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "18")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1D6FD6")
    border.append(bottom)
    rule._p.get_or_add_pPr().append(border)


def add_abstract_box(document: Document, text: str, keywords: str) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    cell = table.cell(0, 0)
    shade_cell(cell, "F3F8FC")
    paragraph = cell.paragraphs[0]
    title_run = paragraph.add_run("Abstract. ")
    title_run.bold = True
    title_run.font.name = "Times New Roman"
    title_run.font.size = Pt(11)
    body_run = paragraph.add_run(text)
    body_run.font.name = "Times New Roman"
    body_run.font.size = Pt(11)
    keyword_paragraph = cell.add_paragraph()
    keyword_label = keyword_paragraph.add_run("Keywords. ")
    keyword_label.bold = True
    keyword_label.font.name = "Times New Roman"
    keyword_label.font.size = Pt(10.5)
    keyword_run = keyword_paragraph.add_run(keywords)
    keyword_run.font.name = "Times New Roman"
    keyword_run.font.size = Pt(10.5)
    document.add_paragraph()


def add_paragraphs(document: Document, paragraphs: list[str]) -> None:
    for text in paragraphs:
        document.add_paragraph(text)


def add_bullet_list(document: Document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(4)
        run = paragraph.add_run(item)
        run.font.name = "Times New Roman"
        run.font.size = Pt(11)


def add_table(document: Document, title: str, headers: list[str], rows: list[list[str]]) -> None:
    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = caption.add_run(title)
    run.bold = True
    run.font.name = "Cambria"
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor.from_string("12304B")

    table = document.add_table(rows=len(rows) + 1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for index, header in enumerate(headers):
        cell = table.cell(0, index)
        shade_cell(cell, "12304B")
        set_cell_text(cell, header, bold=True, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
    for row_index, row in enumerate(rows, start=1):
        for col_index, value in enumerate(row):
            cell = table.cell(row_index, col_index)
            if row_index % 2 == 0:
                shade_cell(cell, "F5F8FB")
            set_cell_text(cell, value, align=WD_ALIGN_PARAGRAPH.CENTER if col_index > 0 else WD_ALIGN_PARAGRAPH.LEFT)
    document.add_paragraph()


def add_figure(document: Document, title: str, path: Path, width: float = 6.0) -> None:
    document.add_picture(str(path), width=Inches(width))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = caption.add_run(title)
    run.font.name = "Times New Roman"
    run.font.size = Pt(9.5)
    run.italic = True
    run.font.color.rgb = RGBColor.from_string("425466")
    document.add_paragraph()


def add_reference_list(document: Document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.18)
        paragraph.paragraph_format.first_line_indent = Inches(-0.18)
        paragraph.paragraph_format.space_after = Pt(4)
        run = paragraph.add_run(item)
        run.font.name = "Times New Roman"
        run.font.size = Pt(10)


def build_tables(
    modelnet: dict[str, dict[str, float | int]],
    params: dict[str, int],
    run_baseline: dict[str, object],
    run_attn: dict[str, object],
    stats: dict[str, dict[str, object]],
) -> dict[str, list[list[str]]]:
    branch_rows = [
        ["master", "Upstream baseline", "Reference PointNet / PointNet++ implementation and published README benchmarks."],
        ["feat/apple_silicon", "Platform hardening", "MPS-safe indexing, Apple Silicon training stability, SE exploration, and early attention trials."],
        ["codex/pointnet2-attn-v2-mac-safe", "Mechanism repair", "Masked local attention, hybrid pooling, stronger cross-level fusion, and the corrected ModelNet40 run."],
        ["codex/3dmnist-attnv2-training", "Cross-dataset validation", "3D MNIST dataloader integration plus run_config.json, metrics.csv, and summary.json artifacts."],
    ]
    architecture_rows = [
        [
            "SSG baseline",
            "Max pooling",
            "None",
            "None",
            "1024 -> 512 -> 256 -> C",
            million(params["baseline"]),
        ],
        [
            "SSG + SE",
            "Max pooling",
            "SE after SA1/SA2",
            "None",
            "1024 -> 512 -> 256 -> 128 -> C",
            million(params["se"]),
        ],
        [
            "Attention v1",
            "Unmasked attention pooling",
            "SE after SA1/SA2",
            "3 mean-pooled tokens",
            "256 -> 128 -> C",
            million(params["attn_v1"]),
        ],
        [
            "Attention v2",
            "Masked hybrid attention + max",
            "SE after SA1/SA2/SA3",
            "Mean/max tokens + CLS + global residual",
            "512 -> 256 -> 128 -> C",
            million(params["attn_v2"]),
        ],
    ]
    modelnet_rows = []
    for key in MODEL_PATHS:
        modelnet_rows.append(
            [
                MODEL_LABELS[key],
                str(modelnet[key]["batch_size"]),
                str(modelnet[key]["best_epoch"]),
                pct(modelnet[key]["best_instance"]),
                pct(modelnet[key]["best_class"]),
                pct(modelnet[key]["final_train"]),
                pct(modelnet[key]["final_test_instance"]),
                pct(modelnet[key]["final_test_class"]),
            ]
        )
    reference_rows = [
        ["PointNet++ paper (XYZ only)", "Official paper [2]", "90.700%"],
        ["PointNet++ paper (XYZ + normals)", "Official paper [2]", "91.900%"],
        ["Repository README SSG (XYZ only)", "Repository README", "92.200%"],
        ["Repository README SSG (XYZ + normals)", "Repository README", "92.400%"],
        ["Local retrain SSG (XYZ only)", "Inspected text log", pct(modelnet["baseline"]["best_instance"])],
    ]
    threed_rows = [
        [
            "SSG baseline",
            str(run_baseline["epochs"]),
            str(run_baseline["batch_size"]),
            pct(run_baseline["best_instance_acc"]),
            pct(run_baseline["best_class_acc"]),
            str(run_baseline["best_epoch"]),
        ],
        [
            "Attention v2",
            str(run_attn["epochs"]),
            str(run_attn["batch_size"]),
            pct(run_attn["best_instance_acc"]),
            pct(run_attn["best_class_acc"]),
            str(run_attn["best_epoch"]),
        ],
    ]
    param_rows = []
    baseline_count = params["baseline"]
    for key in MODEL_PATHS:
        delta = (params[key] - baseline_count) / baseline_count * 100 if key != "baseline" else 0.0
        param_rows.append(
            [
                MODEL_LABELS[key],
                f"{params[key]:,}",
                "baseline" if key == "baseline" else f"{delta:+.2f}%",
            ]
        )
    dataset_rows = [
        ["Train", str(stats["train"]["n"]), f'{stats["train"]["min"]} to {stats["train"]["max"]}', f'{stats["train"]["avg"]:.2f}', "10"],
        ["Test", str(stats["test"]["n"]), f'{stats["test"]["min"]} to {stats["test"]["max"]}', f'{stats["test"]["avg"]:.2f}', "10"],
    ]
    return {
        "branches": branch_rows,
        "architecture": architecture_rows,
        "modelnet": modelnet_rows,
        "reference": reference_rows,
        "3dmnist": threed_rows,
        "params": param_rows,
        "dataset": dataset_rows,
    }


def create_document(
    modelnet: dict[str, dict[str, float | int]],
    params: dict[str, int],
    top_params: list[tuple[str, int]],
    baseline_summary: dict[str, object],
    attn_summary: dict[str, object],
    baseline_metrics: list[dict[str, float]],
    attn_metrics: list[dict[str, float]],
    stats: dict[str, dict[str, object]],
    figures: dict[str, Path],
) -> Document:
    tables = build_tables(modelnet, params, baseline_summary, attn_summary, stats)
    document = Document()
    style_document(document)
    add_title_block(document)

    abstract = (
        "This paper studies how attention should be introduced into PointNet++ classification without "
        "breaking the local invariances and numerical robustness that make the original architecture effective. "
        "The study begins from the PointNet++ single-scale grouping baseline, then follows three "
        "architectural branches: channel recalibration with Squeeze-and-Excitation (SE), an initial "
        "attention design that underperforms, and a repaired attention-v2 model that corrects the observed "
        "failure modes. The central principle is that PointNet++ neighborhood padding is benign for max "
        "pooling but mathematically inconsistent with naive softmax attention, because duplicated padding "
        "points absorb real probability mass. Attention-v2 addresses that mismatch with explicit neighbor "
        "validity masks, hybrid attention-plus-max pooling, richer cross-level tokenization, and a stronger "
        "classifier head, while also adding Apple Silicon / Metal Performance Shaders safeguards for masking, "
        "indexing, and tensor layout. On ModelNet40, the repaired model improves sharply over attention-v1 "
        f"from {pct(modelnet['attn_v1']['best_instance'])} to {pct(modelnet['attn_v2']['best_instance'])} "
        f"instance accuracy and from {pct(modelnet['attn_v1']['best_class'])} to {pct(modelnet['attn_v2']['best_class'])} "
        f"class accuracy, but it still trails the strongest local SSG baseline by {abs(modelnet['baseline']['best_instance'] - modelnet['attn_v2']['best_instance']) * 100:.3f} "
        f"instance points. On 3D MNIST, attention-v2 slightly exceeds the baseline in both best instance "
        f"accuracy ({pct(attn_summary['best_instance_acc'])} vs. {pct(baseline_summary['best_instance_acc'])}) "
        f"and best class accuracy ({pct(attn_summary['best_class_acc'])} vs. {pct(baseline_summary['best_class_acc'])}). "
        "The final conclusion is intentionally balanced: the repaired attention design is technically sound "
        "and cross-dataset viable, but its parameter cost is substantially larger and it does not justify "
        "replacing a strong PointNet++ SSG baseline on ModelNet40."
    )
    keywords = "point cloud classification; PointNet++; masked attention; hybrid local pooling; Apple Silicon; MPS; technical report"
    add_abstract_box(document, abstract, keywords)

    document.add_heading("1. Introduction", level=1)
    add_paragraphs(
        document,
        [
            "Deep learning on point sets must reconcile three constraints at once: permutation invariance, local geometric structure, and computational stability. PointNet established a strong permutation-invariant baseline by combining shared pointwise multilayer perceptrons with a global symmetric max operator [1]. PointNet++ extended that idea with hierarchical set abstraction, radius-based grouping, and local feature aggregation, which remains one of the most durable design templates for point cloud classification [2].",
            "The central question in this paper is not whether attention is fashionable, but whether it is principled in the specific context of PointNet++. PointNet++ relies on fixed-size local groups and max pooling because those choices are robust to irregular neighborhood density and missing neighbors. Any attention mechanism that replaces or augments that path therefore has to preserve two properties at once: it must remain faithful to the real set of neighbors, and it must not destabilize the strong local aggregation bias that already works well.",
            "Two design goals follow from that reasoning. The first is infrastructural: the implementation must remain stable on Apple Silicon and Metal Performance Shaders so that masking, indexing, and pooling behavior do not introduce accidental regressions. The second is architectural: attention should only be retained when it improves the representation for a principled reason, such as better weighting of valid neighbors or better preservation of multi-scale information. This is why the paper treats the full path from baseline to attention v1 to attention v2 as a methodological argument rather than as a sequence of ad hoc model variants.",
        ],
    )
    add_bullet_list(
        document,
        [
            "A principled diagnosis of why naive local softmax attention is incompatible with PointNet++ duplicate-neighbor padding.",
            "A repaired attention-v2 architecture with explicit neighbor masks, hybrid local pooling, richer cross-level tokenization, and a stronger classifier head.",
            "A platform-aware training pipeline for Apple Silicon, plus dataset-aware logging that turns one-off runs into reusable experiment artifacts.",
            "A balanced evaluation across ModelNet40 and 3D MNIST showing stronger class balance and modest cross-dataset gains, but no decisive ModelNet40 top-line win.",
        ],
    )

    document.add_heading("2. Design Motivation and Related Work", level=1)
    add_paragraphs(
        document,
        [
            "Within point cloud learning, PointNet [1] and PointNet++ [2] provide the direct architectural lineage for this work. PointCNN and PointConv explored learnable local operators that more explicitly model neighborhood structure [6, 7], while Dynamic Graph CNN emphasized dynamically updated local graphs for feature extraction [4]. Those families show that gains often emerge from better local aggregation rather than from replacing permutation invariance outright.",
            "Attention mechanisms entered this space through both general permutation-invariant architectures and point-cloud-specific transformers. Set Transformer showed that attention can be used as a principled pooling and set-processing primitive [5], while transformer-based point models such as Point Transformer applied self-attention to 3D neighborhoods more directly [8]. At the same time, Squeeze-and-Excitation networks demonstrated that lightweight channel recalibration can improve representation quality without redesigning the full backbone [3].",
            "The current work occupies a narrower but practically important position. It is not a new state-of-the-art point cloud architecture, nor is it a general survey of transformer methods. Instead, it asks a focused design question: when should attention replace max pooling, and when should it cooperate with max pooling instead? The broader literature on realistic point cloud benchmarks, including ScanObjectNN [9], motivates the decision to treat 3D MNIST as a useful secondary validation rather than a definitive generalization benchmark.",
        ],
    )
    add_table(
        document,
        "Table 1. Project evolution and experimental branch roles.",
        ["Branch", "Role", "Main contribution"],
        tables["branches"],
    )
    add_figure(
        document,
        "Figure 1. The project evolves from the upstream baseline through Apple Silicon hardening, attention-v2 repair, and cross-dataset logging support.",
        figures["timeline"],
    )

    document.add_heading("3. Methodology", level=1)
    document.add_heading("3.1 Baseline PointNet++ SSG", level=2)
    add_paragraphs(
        document,
        [
            "The baseline classifier follows the standard PointNet++ single-scale grouping design. It applies three set abstraction stages: SA1 samples 512 centroids with radius 0.2 and 32 neighbors, SA2 samples 128 centroids with radius 0.4 and 64 neighbors, and SA3 performs global aggregation into a 1024-dimensional descriptor. The classifier head maps 1024 -> 512 -> 256 -> num_class with ReLU, batch normalization, and dropout. This baseline is both historically grounded and locally strong, which raises the bar for any augmentation.",
            "Crucially, the local aggregation in this baseline is max pooling. That operator is robust to duplicated neighbors because repeating a value does not change the maximum. The core design problem for the rest of the paper is therefore not simply how to add attention, but how to add it without discarding the robustness that max pooling already gives for free.",
        ],
    )
    document.add_heading("3.2 SE and Attention v1", level=2)
    add_paragraphs(
        document,
        [
            "The SE variant keeps the PointNet++ feature hierarchy intact but inserts squeeze-and-excitation blocks after SA1 and SA2. It also deepens the classifier head to 1024 -> 512 -> 256 -> 128 -> num_class. This model tests a conservative hypothesis: stronger channel weighting may improve class discrimination while preserving PointNet++ locality.",
            "Attention v1 is more ambitious. It replaces local max pooling with learned attention weights, retains SE channel recalibration, and introduces cross-level fusion by reducing l1, l2, and l3 into three tokens processed by multi-head self-attention. In intent, the model aims to preserve finer local structure than max pooling and to exploit multi-scale information that the global l3 descriptor alone might miss.",
        ],
    )
    document.add_heading("3.3 Why Attention v1 Fails", level=2)
    add_paragraphs(
        document,
        [
            "The strongest failure mode is a mismatch between PointNet++ grouping semantics and softmax attention. In the shared utility path, query_ball_point pads missing neighbors by duplicating the first valid point so that downstream tensor shapes remain fixed. That duplication is benign for max pooling but not for attention: the copied entries receive real probability mass and bias the neighborhood aggregation. In short, the model is not attending over the true set of neighbors; it is attending over a padded proxy whose measure has been changed.",
            "A second issue is excessive compression before fusion. Attention v1 collapses each abstraction level into a single mean-pooled token and then averages the attended tokens again, which removes much of the structure that cross-level attention was meant to preserve. A third issue lies in the classifier. The post-fusion head does not introduce enough nonlinear depth, so the richer features are not followed by a suitably expressive decision boundary.",
        ],
    )
    add_figure(
        document,
        "Figure 2. Attention v2 is best interpreted as a repair program for concrete attention-v1 failure modes rather than as a minor extension.",
        figures["repair"],
    )
    document.add_heading("3.4 Attention v2", level=2)
    add_paragraphs(
        document,
        [
            "Attention v2 corrects the three architectural problems and adds platform-aware numerical safeguards. First, query_ball_point_with_mask and sample_and_group_with_mask return both gathered indices and a validity mask. Softmax is applied only over valid neighbors, while padded slots are used solely for safe indexing. This aligns the implementation with the mathematical assumption that attention should normalize over real elements only.",
            "Second, local pooling becomes hybrid rather than purely attentional. Each neighborhood produces both attention-pooled features and a max-pooled residual, and a small fusion block recombines them. This design preserves the strong inductive bias of PointNet / PointNet++ max aggregation while allowing learned reweighting to contribute when it is useful.",
            "Third, cross-level fusion is expanded substantially. Attention v2 uses mean and max summaries for l1 and l2, adds a learnable CLS token, and preserves a direct residual path from the global l3 representation. The final classifier is widened to 512 -> 256 -> 128 -> num_class. Together, these changes make the final stage more expressive and reduce the risk that useful multi-scale signals are averaged away too early.",
        ],
    )
    add_table(
        document,
        "Table 2. Architectural comparison across the four studied classifiers.",
        ["Model", "Local aggregation", "Channel attention", "Cross-level fusion", "Classifier head", "Parameters"],
        tables["architecture"],
    )

    document.add_heading("4. Implementation and Experimental Protocol", level=1)
    add_paragraphs(
        document,
        [
            "Engineering support is a first-class part of the argument because attention is only meaningful if the execution path is stable. device_utils.py adds automatic device selection across MPS, CUDA, and CPU. models/pointnet2_utils.py adds MPS-safe index clamping, torch.where-based masking in place of boolean assignment, and contiguous permutations. Attention v2 adds further protections such as finite large-negative masking values instead of -inf and torch.nan_to_num after multi-head attention. These changes are especially important on Apple Silicon, where indexing and masking behavior is less forgiving than on CUDA.",
            "The training pipeline was upgraded so that experiments preserve their own provenance. train_classification.py writes run_config.json, metrics.csv, and summary.json, copies the exact model and loader files used during training, and supports both ModelNet and 3D MNIST through a common interface. test_classification.py reads run_config.json automatically, which reduces configuration drift at evaluation time.",
            "ModelNet40 remains the primary benchmark because it is the dataset most directly tied to the PointNet++ literature [2]. However, the older ModelNet40 branches still record results mainly in text logs, and the batch size was not fully controlled across all four variants: the baseline and attention v1 use batch size 24, while SE and attention v2 use batch size 16. This does not invalidate the comparison, but it does weaken any claim that the ModelNet40 ranking is perfectly controlled.",
            "3D MNIST was integrated as a second classification dataset with HDF5-backed train and test splits. The local copy contains 5,000 training samples and 1,000 test samples across 10 classes. Point counts vary substantially before subsampling, so the dataset also serves as a useful stress test for the dataloader and for the point sampling and normalization path. For HDF5 safety on macOS, effective dataloader workers are forced to zero during these runs.",
        ],
    )
    add_table(
        document,
        "Table 3. Local 3D MNIST dataset characteristics from the HDF5 train and test splits.",
        ["Split", "Samples", "Point-count range", "Mean raw points", "Classes"],
        tables["dataset"],
    )
    add_figure(
        document,
        "Figure 3. The local 3D MNIST copy is reasonably class-balanced and exhibits a wide raw point-count range before subsampling to 1,024 points.",
        figures["dataset"],
    )

    document.add_heading("5. Results", level=1)
    document.add_heading("5.1 ModelNet40", level=2)
    add_paragraphs(
        document,
        [
            f"The strongest local baseline reaches {pct(modelnet['baseline']['best_instance'])} instance accuracy and {pct(modelnet['baseline']['best_class'])} class accuracy. SE trades a small instance-accuracy drop for a class-accuracy gain, which is consistent with channel recalibration helping balance predictions. Attention v1 performs substantially worse, confirming that its issues are material rather than cosmetic.",
            f"Attention v2 recovers much of that damage. Relative to attention v1, it gains {(modelnet['attn_v2']['best_instance'] - modelnet['attn_v1']['best_instance']) * 100:.3f} instance points and {(modelnet['attn_v2']['best_class'] - modelnet['attn_v1']['best_class']) * 100:.3f} class-accuracy points. Relative to the strongest baseline, however, it remains {(modelnet['baseline']['best_instance'] - modelnet['attn_v2']['best_instance']) * 100:.3f} points lower on the main top-line metric even while exceeding the baseline by {(modelnet['attn_v2']['best_class'] - modelnet['baseline']['best_class']) * 100:.3f} class-accuracy points.",
            "This pattern matters because it clarifies what the repaired attention is actually buying. The gain is not a universal improvement in top-line accuracy; instead, the model appears to redistribute representational capacity toward more balanced class behavior. That is a principled outcome for a mechanism designed to weight valid neighbors and preserve cross-level context, even if it is not yet enough to displace a strong PointNet++ SSG baseline on the primary benchmark.",
        ],
    )
    add_table(
        document,
        "Table 4. ModelNet40 outcomes extracted directly from the four text log files.",
        ["Model", "Batch size", "Best epoch", "Best inst.", "Best class", "Final train inst.", "Final test inst.", "Final test class"],
        tables["modelnet"],
    )
    add_figure(
        document,
        "Figure 4. Attention v2 closes most of the gap opened by attention v1, but the baseline retains the highest ModelNet40 instance accuracy.",
        figures["modelnet"],
    )
    add_table(
        document,
        "Table 5. Reference accuracy context for PointNet++ on ModelNet40.",
        ["Method", "Source", "Instance accuracy"],
        tables["reference"],
    )
    document.add_heading("5.2 3D MNIST", level=2)
    add_paragraphs(
        document,
        [
            f"3D MNIST provides the cleanest head-to-head comparison in this study because the baseline and attention-v2 runs share the same batch size, epoch count, optimizer, learning rate, weight decay, point count, and device. Under that matched setup, attention v2 achieves {pct(attn_summary['best_instance_acc'])} best instance accuracy and {pct(attn_summary['best_class_acc'])} best class accuracy, slightly exceeding the baseline values of {pct(baseline_summary['best_instance_acc'])} and {pct(baseline_summary['best_class_acc'])}, respectively.",
            "The convergence curves are informative. The baseline learns much faster in the first few epochs, which is expected for a simpler architecture with fewer parameters. Attention v2 starts slowly, then catches up and eventually produces stronger best metrics later in training. This delayed payoff is consistent with attention v2 being a higher-capacity model that needs longer to stabilize, but which can eventually refine class balance more effectively once optimization catches up.",
            "The gain is modest and should not be overstated. 3D MNIST is simpler than realistic object benchmarks and is not a canonical PointNet++ dataset. Still, the result matters because it shows that the repaired attention design is not inherently broken; under a matched setup on a second dataset, it can edge out the baseline.",
        ],
    )
    add_table(
        document,
        "Table 6. Matched 3D MNIST comparison using the structured experiment artifacts.",
        ["Model", "Epochs", "Batch size", "Best inst.", "Best class", "Best epoch"],
        tables["3dmnist"],
    )
    add_figure(
        document,
        "Figure 5. On 3D MNIST, attention v2 converges more slowly but eventually reaches slightly stronger best metrics than the baseline.",
        figures["curves"],
    )
    add_table(
        document,
        "Table 7. Parameter cost across the four classifiers.",
        ["Model", "Parameters", "Delta vs. baseline"],
        tables["params"],
    )
    add_figure(
        document,
        "Figure 6. The repaired attention stack trades efficiency for robustness and richer fusion, with attention v2 reaching 5.95M parameters.",
        figures["efficiency"],
    )

    document.add_heading("6. Discussion", level=1)
    largest_name, largest_size = top_params[0]
    add_paragraphs(
        document,
        [
            "The final model's strongest contribution is methodological rather than purely numerical. It identifies a subtle but consequential incompatibility between PointNet++ neighborhood padding and naive softmax attention, then fixes that incompatibility with an explicit mask path. That diagnosis is portable: any point-cloud pipeline that mixes fixed-size neighborhood tensors with softmax-based local pooling should be audited for the same issue.",
            f"The main drawback is efficiency. Attention v2 contains {params['attn_v2']:,} parameters, or {(params['attn_v2'] - params['baseline']) / params['baseline'] * 100:.2f}% more than the baseline. The single largest block is {largest_name}, which alone contains {largest_size:,} parameters. This confirms that the current design should be treated as an accuracy-oriented prototype rather than an optimized deployment model.",
            "The results also show the limits of local improvements in a mature architecture family. A strong PointNet++ implementation already exceeds the original paper's XYZ-only number, so the challenge is not merely to beat the 2017 benchmark but to beat a strong retrained baseline. Within that harder comparison, attention v2 improves class balance and rescues the failed attention branch, but it does not earn a clear replacement case on ModelNet40.",
        ],
    )
    document.add_heading("7. Limitations and Threats to Validity", level=1)
    add_bullet_list(
        document,
        [
            "ModelNet40 runs are not fully controlled because batch size differs across variants and the older branches record results mainly in text logs.",
            "No multi-seed evaluation was performed, so the reported rankings may still contain seed sensitivity.",
            "No ablation isolates masked attention, hybrid pooling, cross-level fusion, or the stronger classifier independently.",
            "3D MNIST is a useful cross-dataset check, but it is easier and less realistic than benchmarks such as ScanObjectNN [9].",
        ],
    )
    document.add_heading("8. Conclusion", level=1)
    add_paragraphs(
        document,
        [
            "This study shows that attention can improve a PointNet++ SSG classifier only when the surrounding assumptions are repaired as carefully as the attention block itself. Attention v1 fails because it violates the semantics of padded neighborhoods and compresses too much information before fusion. Attention v2 succeeds in fixing those issues with explicit masking, hybrid pooling, richer tokenization, and a stronger classifier, while also making the implementation safe on Apple Silicon.",
            f"The repaired model is clearly better than attention v1, reaches {pct(attn_summary['best_instance_acc'])} / {pct(attn_summary['best_class_acc'])} on 3D MNIST, and improves ModelNet40 class accuracy relative to the baseline. Yet the baseline still holds the best ModelNet40 instance accuracy at {pct(modelnet['baseline']['best_instance'])}, and it does so with far fewer parameters. The most defensible conclusion is therefore balanced: the attention redesign is technically worthwhile and empirically credible, but it is not yet the best default replacement for PointNet++ SSG on the primary benchmark.",
        ],
    )
    document.add_heading("References", level=1)
    add_reference_list(document, REFERENCE_TEXT)
    return document


def main() -> None:
    ensure_dirs()

    modelnet = {
        "baseline": parse_modelnet_log(ROOT / "log" / "classification" / "pointnet2_cls_ssg" / "logs" / "pointnet2_cls_ssg.txt"),
        "se": parse_modelnet_log(ROOT / "log" / "classification" / "se" / "logs" / "pointnet2_cls_ssg_se.txt"),
        "attn_v1": parse_modelnet_log(ROOT / "log" / "classification" / "attn" / "logs" / "pointnet2_cls_ssg_attn.txt"),
        "attn_v2": parse_modelnet_log(ROOT / "log" / "classification" / "attnv2" / "logs" / "pointnet2_cls_ssg_attn_v2.txt"),
    }
    baseline_summary = load_run_summary(ROOT / "log" / "classification" / "3dmnist_pointnet2_ssg" / "summary.json")
    attn_summary = load_run_summary(ROOT / "log" / "classification" / "3dmnist_attnv2" / "summary.json")
    baseline_metrics = load_metrics(ROOT / "log" / "classification" / "3dmnist_pointnet2_ssg" / "metrics.csv")
    attn_metrics = load_metrics(ROOT / "log" / "classification" / "3dmnist_attnv2" / "metrics.csv")
    params, top_params = parameter_counts()
    stats = dataset_stats()
    git_data = git_overview()

    figures = {
        "timeline": ASSET_DIR / "timeline.png",
        "repair": ASSET_DIR / "repair_map.png",
        "modelnet": ASSET_DIR / "modelnet_results.png",
        "curves": ASSET_DIR / "3dmnist_curves.png",
        "efficiency": ASSET_DIR / "efficiency.png",
        "dataset": ASSET_DIR / "dataset_stats.png",
    }
    save_timeline_figure(figures["timeline"])
    save_repair_figure(figures["repair"])
    save_modelnet_figure(figures["modelnet"], modelnet)
    save_3dmnist_curves(figures["curves"], baseline_metrics, attn_metrics)
    save_efficiency_figure(figures["efficiency"], params, modelnet)
    save_dataset_figure(figures["dataset"], stats)

    document = create_document(
        modelnet,
        params,
        top_params,
        baseline_summary,
        attn_summary,
        baseline_metrics,
        attn_metrics,
        stats,
        figures,
    )
    core_props = document.core_properties
    core_props.author = "OpenAI Codex"
    core_props.title = "Repairing Attention for PointNet++ Classification"
    core_props.subject = "Single-column technical paper"
    core_props.keywords = "point cloud, PointNet++, attention, MPS, 3D MNIST"
    core_props.comments = "Generated from code and experiment artifacts."
    document.save(DOCX_PATH)

    print(f"Wrote {DOCX_PATH}")
    print(f"Assets directory: {ASSET_DIR}")
    print(f"Available branches: {', '.join(git_data['branches'])}")
    print(f"Top attention-v2 parameter block: {top_params[0][0]} = {top_params[0][1]:,}")


if __name__ == "__main__":
    main()
